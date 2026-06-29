"""Unit tests for Markdown → Word conversion."""
import pytest
import tempfile
import os
from io import BytesIO
from pathlib import Path


def _write_md(content: str) -> str:
    """将字符串写入临时 .md 文件，返回路径。"""
    tmp = tempfile.NamedTemporaryFile(suffix=".md", delete=False, mode="w", encoding="utf-8")
    tmp.write(content)
    tmp.close()
    return tmp.name


class TestConvertMdToDocxBytes:
    """convert_md_to_docx_bytes 函数测试。"""

    def test_basic_markdown_produces_valid_docx(self):
        """基本 Markdown 生成有效 .docx（ZIP magic bytes）。"""
        from convert_to_docx import convert_md_to_docx_bytes

        md_path = _write_md("# 测试标题\n\n测试正文内容。\n\n- 列表项1\n- 列表项2")
        try:
            buf = convert_md_to_docx_bytes(md_path)
            magic = buf.read(4)
            assert magic == b'PK\x03\x04', f"Expected ZIP magic, got {magic!r}"
            buf.close()
        finally:
            os.unlink(md_path)

    def test_output_can_be_opened_by_python_docx(self):
        """生成的 BytesIO 可被 python-docx 打开并读取段落。"""
        from convert_to_docx import convert_md_to_docx_bytes
        from docx import Document

        md_path = _write_md("# H1\n\n段落一。\n\n## H2\n\n段落二。")
        try:
            buf = convert_md_to_docx_bytes(md_path)
            tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
            tmp.write(buf.read())
            tmp.close()
            buf.close()

            doc = Document(tmp.name)
            # 应至少包含标题和正文段落
            assert len(doc.paragraphs) >= 3, f"Expected >=3 paragraphs, got {len(doc.paragraphs)}"
            os.unlink(tmp.name)
        finally:
            os.unlink(md_path)

    def test_empty_content_does_not_crash(self):
        """空内容不崩溃，生成有效 docx。"""
        from convert_to_docx import convert_md_to_docx_bytes

        md_path = _write_md("")
        try:
            buf = convert_md_to_docx_bytes(md_path)
            assert buf.read(4) == b'PK\x03\x04'
            buf.close()
        finally:
            os.unlink(md_path)

    def test_table_parsing(self):
        """Markdown 表格正确转换为 Word 表格。"""
        from convert_to_docx import convert_md_to_docx_bytes
        from docx import Document

        md_path = _write_md("# 表格测试\n\n| 列A | 列B |\n|-----|-----|\n| a1 | b1 |\n| a2 | b2 |")
        try:
            buf = convert_md_to_docx_bytes(md_path)
            tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
            tmp.write(buf.read())
            tmp.close()
            buf.close()

            doc = Document(tmp.name)
            tables = doc.tables
            assert len(tables) >= 1, f"Expected >=1 table, got {len(tables)}"
            os.unlink(tmp.name)
        finally:
            os.unlink(md_path)

    def test_code_block(self):
        """代码块正确渲染。"""
        from convert_to_docx import convert_md_to_docx_bytes
        from docx import Document

        md_path = _write_md("# 代码测试\n\n```\nprint('hello')\n```\n\n正文。")
        try:
            buf = convert_md_to_docx_bytes(md_path)
            tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
            tmp.write(buf.read())
            tmp.close()
            buf.close()

            doc = Document(tmp.name)
            # 代码行应在段落中
            code_found = any("print" in p.text for p in doc.paragraphs)
            assert code_found, "Code line not found in output"
            os.unlink(tmp.name)
        finally:
            os.unlink(md_path)

    def test_missing_file_raises(self):
        """不存在的文件路径应抛出 FileNotFoundError。"""
        from convert_to_docx import convert_md_to_docx_bytes

        with pytest.raises(FileNotFoundError):
            convert_md_to_docx_bytes("/nonexistent/path/report.md")
